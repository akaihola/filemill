import sys
import pytest
from pathlib import Path
from unittest.mock import MagicMock

from pykofinder.preview import (
    render_preview,
    _preview_md,
    _preview_docx,
    _preview_pptx,
    _preview_pdf,
    _preview_image,
    _preview_desktop,
)


# ── helpers ──────────────────────────────────────────────────────────────────


def _make_pptx(
    path: Path, text: str = "Slide text", bold: bool = False, italic: bool = False
):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    p = txBox.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.bold = bold
    run.font.italic = italic
    prs.save(str(path))


def _fake_mammoth(monkeypatch, html_value: str = "<p>Hello</p>", raises=None):
    """Patch sys.modules['mammoth'] with a mock. Call before _preview_docx."""
    m = MagicMock()
    if raises:
        m.convert_to_html.side_effect = raises
    else:
        m.convert_to_html.return_value = MagicMock(value=html_value)
    monkeypatch.setitem(sys.modules, "mammoth", m)
    return m


# ── render_preview dispatch ───────────────────────────────────────────────────


def test_dispatch_desktop(tmp_path):
    f = tmp_path / "link.desktop"
    f.write_text("[Desktop Entry]\nType=Link\nURL=https://ex.com\nName=Ex\n")
    assert "preview-desktop-link" in render_preview(f)


def test_dispatch_md(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("# Hi")
    assert "preview-md" in render_preview(f)


def test_dispatch_docx(tmp_path, monkeypatch):
    f = tmp_path / "doc.docx"
    f.write_bytes(b"fake")
    _fake_mammoth(monkeypatch)
    assert "preview-docx" in render_preview(f)


def test_dispatch_pptx(tmp_path):
    f = tmp_path / "deck.pptx"
    _make_pptx(f, text="Hello Slide")
    assert "preview-pptx" in render_preview(f)


def test_dispatch_pdf(tmp_path):
    f = tmp_path / "doc.pdf"
    f.write_bytes(b"%PDF")
    assert "preview-pdf" in render_preview(f)


@pytest.mark.parametrize("ext", [".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"])
def test_dispatch_image(tmp_path, ext):
    f = tmp_path / f"img{ext}"
    f.write_bytes(b"fake")
    assert "preview-image" in render_preview(f)


def test_dispatch_unknown_ext(tmp_path):
    """Binary file with unknown extension → preview-unsupported with extension name."""
    f = tmp_path / "data.xyz"
    f.write_bytes(b"\x00\x01\x02\xff\xfe")  # binary; UTF-8 decode fails
    result = render_preview(f)
    assert "preview-unsupported" in result
    assert ".xyz" in result


def test_dispatch_no_ext(tmp_path):
    """Binary file with no extension → '(no extension)' in unsupported message."""
    f = tmp_path / "noext"
    f.write_bytes(b"\x00\x01\x02\xff\xfe")  # binary; UTF-8 decode fails
    result = render_preview(f)
    assert "preview-unsupported" in result
    assert "(no extension)" in result


# ── _preview_md ───────────────────────────────────────────────────────────────


def test_preview_md_valid(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("# Title\n**bold**\n")
    html = _preview_md(f)
    assert "preview-md" in html
    assert "<h1" in html  # anchors_plugin adds id= attribute to headings


def test_preview_md_exception(tmp_path, monkeypatch):
    f = tmp_path / "note.md"
    f.write_text("hi")
    from pykofinder import rendering

    def boom(text):
        raise RuntimeError("render failure")

    monkeypatch.setattr(rendering.md, "render", boom)
    assert "preview-error" in _preview_md(f)


# ── _preview_docx ─────────────────────────────────────────────────────────────


def test_preview_docx_happy_path(tmp_path, monkeypatch):
    f = tmp_path / "doc.docx"
    f.write_bytes(b"fake")
    _fake_mammoth(monkeypatch, html_value="<p>Hello docx</p>")
    html = _preview_docx(f)
    assert "preview-docx" in html
    assert "<p>Hello docx</p>" in html


def test_preview_docx_exception(tmp_path, monkeypatch):
    f = tmp_path / "doc.docx"
    f.write_bytes(b"fake")
    _fake_mammoth(monkeypatch, raises=RuntimeError("bad docx"))
    assert "preview-error" in _preview_docx(f)


# ── _preview_pptx ─────────────────────────────────────────────────────────────


def test_preview_pptx_valid(tmp_path):
    f = tmp_path / "deck.pptx"
    _make_pptx(f, text="Hello Slide")
    html = _preview_pptx(f)
    assert "preview-pptx" in html
    assert "Hello Slide" in html


def test_preview_pptx_bold(tmp_path):
    f = tmp_path / "deck.pptx"
    _make_pptx(f, text="Bold text", bold=True)
    assert "<strong>Bold text</strong>" in _preview_pptx(f)


def test_preview_pptx_italic(tmp_path):
    f = tmp_path / "deck.pptx"
    _make_pptx(f, text="Italic text", italic=True)
    assert "<em>Italic text</em>" in _preview_pptx(f)


def test_preview_pptx_exception(tmp_path):
    """Non-zip bytes cause Presentation() to raise; must return preview-error."""
    f = tmp_path / "deck.pptx"
    f.write_bytes(b"not a zip")
    assert "preview-error" in _preview_pptx(f)


def test_preview_pptx_shape_without_text_frame_is_skipped(tmp_path):
    """A picture shape (has_text_frame=False) must not crash – it is skipped."""
    import io
    from pptx import Presentation
    from pptx.util import Inches

    try:
        from PIL import Image

        img = Image.new("RGB", (1, 1), color="red")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(buf, Inches(0), Inches(0), Inches(1), Inches(1))
        prs.save(str(tmp_path / "deck.pptx"))
    except ImportError:
        pytest.skip("Pillow not installed")
    assert "preview-pptx" in _preview_pptx(tmp_path / "deck.pptx")


def test_preview_pptx_empty_paragraph_is_skipped(tmp_path):
    """A paragraph with no runs produces an empty line_parts list (if branch=False)."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    # First paragraph: has a run with text (covered already)
    txBox.text_frame.paragraphs[0].add_run().text = "visible"
    # Second paragraph: no runs → empty line_parts → if branch is False
    txBox.text_frame.add_paragraph()  # empty paragraph, no runs
    prs.save(str(tmp_path / "deck.pptx"))
    html = _preview_pptx(tmp_path / "deck.pptx")
    assert "visible" in html


# ── _preview_pdf ──────────────────────────────────────────────────────────────


def test_preview_pdf_contains_iframe(tmp_path):
    f = tmp_path / "doc.pdf"
    f.write_bytes(b"%PDF")
    html = _preview_pdf(f)
    assert "preview-pdf" in html
    assert "<iframe" in html
    assert "/raw?path=" in html


# ── _preview_image ────────────────────────────────────────────────────────────


def test_preview_image_contains_img_tag(tmp_path):
    f = tmp_path / "img.png"
    f.write_bytes(b"\x89PNG")
    html = _preview_image(f)
    assert "preview-image" in html
    assert "<img" in html
    assert "/raw?path=" in html


def test_preview_image_alt_is_filename(tmp_path):
    f = tmp_path / "photo.jpg"
    f.write_bytes(b"JFIF")
    assert 'alt="photo.jpg"' in _preview_image(f)


# ── _preview_desktop ──────────────────────────────────────────────────────────


def test_preview_desktop_valid_link(tmp_path):
    f = tmp_path / "link.desktop"
    f.write_text("[Desktop Entry]\nType=Link\nURL=https://example.com\nName=Example\n")
    html = _preview_desktop(f)
    assert "preview-desktop-link" in html
    assert "https://example.com" in html
    assert "Open \u2192" in html  # "Open →"


def test_preview_desktop_http_icon_rendered(tmp_path):
    f = tmp_path / "link.desktop"
    f.write_text(
        "[Desktop Entry]\nType=Link\nURL=https://ex.com\nName=Ex\n"
        "Icon=https://ex.com/icon.png\n"
    )
    assert '<img src="https://ex.com/icon.png"' in _preview_desktop(f)


def test_preview_desktop_local_icon_not_rendered(tmp_path):
    """Non-http icon names must NOT produce an <img> element."""
    f = tmp_path / "link.desktop"
    f.write_text(
        "[Desktop Entry]\nType=Link\nURL=https://ex.com\nName=Ex\n"
        "Icon=application-icon\n"
    )
    assert 'src="application-icon"' not in _preview_desktop(f)


def test_preview_desktop_with_comment(tmp_path):
    f = tmp_path / "link.desktop"
    f.write_text(
        "[Desktop Entry]\nType=Link\nURL=https://ex.com\nName=Ex\nComment=My note\n"
    )
    assert "My note" in _preview_desktop(f)


def test_preview_desktop_no_comment_no_paragraph(tmp_path):
    """Absent Comment key must not produce the grey comment <p>."""
    f = tmp_path / "link.desktop"
    f.write_text("[Desktop Entry]\nType=Link\nURL=https://ex.com\nName=Ex\n")
    assert "color:#aaa" not in _preview_desktop(f)


def test_preview_desktop_non_link_type(tmp_path):
    f = tmp_path / "app.desktop"
    f.write_text("[Desktop Entry]\nType=Application\nName=App\n")
    html = _preview_desktop(f)
    assert "preview-unsupported" in html
    assert "no URL to open" in html


def test_preview_desktop_link_type_empty_url(tmp_path):
    """Type=Link but no URL key → 'no URL to open' message."""
    f = tmp_path / "link.desktop"
    f.write_text("[Desktop Entry]\nType=Link\nName=NoURL\n")
    html = _preview_desktop(f)
    assert "preview-unsupported" in html
    assert "no URL to open" in html


def test_preview_desktop_no_section(tmp_path):
    """File without [Desktop Entry] section → 'Not a valid .desktop file'."""
    f = tmp_path / "bad.desktop"
    f.write_text("[OtherSection]\nFoo=Bar\n")
    assert "Not a valid .desktop file" in _preview_desktop(f)


def test_preview_desktop_exception(tmp_path, monkeypatch):
    """ConfigParser.read raises → preview-error div."""
    f = tmp_path / "link.desktop"
    f.write_text("irrelevant")
    import configparser as cp_mod

    def boom(self, *a, **kw):
        raise RuntimeError("parse failure")

    monkeypatch.setattr(cp_mod.ConfigParser, "read", boom)
    assert "preview-error" in _preview_desktop(f)


# ── #7 raw UTF-8 text fallback ────────────────────────────────────────────────


def test_raw_utf8_text_file_shows_preview_raw(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("hello world\nline 2\n")
    result = render_preview(f)
    assert "preview-raw" in result
    assert "hello world" in result


def test_binary_file_shows_unsupported(tmp_path):
    f = tmp_path / "data.bin"
    f.write_bytes(b"\x00\x01\x02\xff\xfe")
    result = render_preview(f)
    assert "No preview available" in result


def test_large_text_file_shows_unsupported(tmp_path):
    f = tmp_path / "big.log"
    f.write_bytes(b"x" * (256 * 1024 + 1))
    result = render_preview(f)
    assert "No preview available" in result


def test_raw_text_html_escaped(tmp_path):
    f = tmp_path / "code.txt"
    f.write_text("<script>alert('xss')</script>\n")
    result = render_preview(f)
    assert "<script>" not in result
    assert "&lt;script&gt;" in result


# ── #14 syntax highlighting ───────────────────────────────────────────────────


def test_python_file_is_syntax_highlighted(tmp_path):
    f = tmp_path / "hello.py"
    f.write_text("def hello():\n    return 'world'\n")
    result = render_preview(f)
    assert "preview-code" in result
    assert "highlight" in result  # Pygments wraps in div.highlight


def test_js_file_is_syntax_highlighted(tmp_path):
    f = tmp_path / "app.js"
    f.write_text("function hello() { return 'world'; }\n")
    result = render_preview(f)
    assert "preview-code" in result


def test_large_code_file_falls_through(tmp_path):
    """Files over 512 KB must not be syntax-highlighted (fallback to raw or unsupported)."""
    f = tmp_path / "big.py"
    f.write_bytes(b"# comment\n" * 52429)  # >512 KB
    result = render_preview(f)
    # Must NOT have preview-code (too large for highlighting)
    assert "preview-code" not in result


def test_binary_py_file_falls_through(tmp_path):
    """A .py file with non-UTF8 bytes must not crash – fall through to next handler."""
    f = tmp_path / "bytes.py"
    f.write_bytes(b"\xff\xfe invalid utf-8")
    result = render_preview(f)  # must not raise
    assert "preview-code" not in result
