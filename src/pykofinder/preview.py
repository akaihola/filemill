import configparser
import html as html_lib
from pathlib import Path
from urllib.parse import quote as urlquote


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}


def render_preview(path: Path) -> str:
    """Return an HTML string (inner body fragment) for the given file path."""
    ext = path.suffix.lower()

    if ext == ".desktop":
        return _preview_desktop(path)
    elif ext == ".md":
        return _preview_md(path)
    elif ext == ".docx":
        return _preview_docx(path)
    elif ext == ".pptx":
        return _preview_pptx(path)
    elif ext == ".pdf":
        return _preview_pdf(path)
    elif ext in IMAGE_EXTS:
        return _preview_image(path)
    else:
        safe_ext = html_lib.escape(ext or "(no extension)")
        return f'<div class="preview-unsupported"><em>No preview available for {safe_ext} files.</em></div>'


def _preview_md(path: Path) -> str:
    try:
        from pykofinder.rendering import md as md_renderer

        html_body = md_renderer.render(path.read_text(encoding="utf-8"))
        return f'<div class="preview-md">{html_body}</div>'
    except Exception as e:
        return (
            f'<div class="preview-error">Preview error: {html_lib.escape(str(e))}</div>'
        )


def _preview_docx(path: Path) -> str:
    try:
        import mammoth

        with open(path, "rb") as f:
            result = mammoth.convert_to_html(f)
        return f'<div class="preview-docx">{result.value}</div>'
    except Exception as e:
        return (
            f'<div class="preview-error">Preview error: {html_lib.escape(str(e))}</div>'
        )


def _preview_pptx(path: Path) -> str:
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        slides_html = []
        for i, slide in enumerate(prs.slides, 1):
            parts = [f'<div class="slide"><h3>Slide {i}</h3>']
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    line_parts = []
                    for run in para.runs:
                        text = html_lib.escape(run.text)
                        if run.font.bold:
                            text = f"<strong>{text}</strong>"
                        if run.font.italic:
                            text = f"<em>{text}</em>"
                        line_parts.append(text)
                    if line_parts:
                        parts.append(f"<p>{''.join(line_parts)}</p>")
            parts.append("</div>")
            slides_html.append("".join(parts))
        return f'<div class="preview-pptx">{"".join(slides_html)}</div>'
    except Exception as e:
        return (
            f'<div class="preview-error">Preview error: {html_lib.escape(str(e))}</div>'
        )


def _preview_pdf(path: Path) -> str:
    try:
        src = f"/raw?path={urlquote(str(path))}"
        return f'<div class="preview-pdf"><iframe src="{src}"></iframe></div>'
    except Exception as e:
        return (
            f'<div class="preview-error">Preview error: {html_lib.escape(str(e))}</div>'
        )


def _preview_image(path: Path) -> str:
    try:
        src = f"/raw?path={urlquote(str(path))}"
        safe_name = html_lib.escape(path.name)
        return f'<div class="preview-image"><img src="{src}" alt="{safe_name}"></div>'
    except Exception as e:
        return (
            f'<div class="preview-error">Preview error: {html_lib.escape(str(e))}</div>'
        )


def _preview_desktop(path: Path) -> str:
    """Preview a .desktop link file — show name, icon, and a clickable URL."""
    try:
        cp = configparser.ConfigParser(interpolation=None)
        cp.read(str(path), encoding="utf-8")
        if "Desktop Entry" not in cp:
            return '<div class="preview-unsupported"><em>Not a valid .desktop file.</em></div>'
        entry = cp["Desktop Entry"]
        entry_type = entry.get("Type", "").strip()
        name = html_lib.escape(entry.get("Name", path.stem).strip())
        url = entry.get("URL", "").strip()
        icon = entry.get("Icon", "").strip()
        comment = html_lib.escape(entry.get("Comment", "").strip())

        if entry_type != "Link" or not url:
            return f'<div class="preview-unsupported"><em>Desktop entry type: {html_lib.escape(entry_type or "unknown")} — no URL to open.</em></div>'

        safe_url = html_lib.escape(url)
        open_href = f"/open-link?path={urlquote(str(path))}"

        icon_html = ""
        if icon.startswith(("http://", "https://")):
            safe_icon = html_lib.escape(icon)
            icon_html = f'<img src="{safe_icon}" alt="" style="width:48px;height:48px;object-fit:contain;flex-shrink:0;">'

        comment_html = (
            f'<p style="margin:0;color:#aaa;font-size:0.9rem;">{comment}</p>'
            if comment
            else ""
        )

        return f"""<div class="preview-desktop-link" style="padding:2rem;display:flex;flex-direction:column;gap:1.2rem;">
  <div style="display:flex;align-items:center;gap:0.8rem;">
    {icon_html}<h2 style="margin:0;font-size:1.4rem;">{name}</h2>
  </div>
  {comment_html}
  <div style="word-break:break-all;">
    🔗 <a href="{safe_url}" target="_blank" rel="noopener noreferrer"
          style="color:#4a9eff;">{safe_url}</a>
  </div>
  <div>
    <a href="{open_href}" target="_blank" rel="noopener noreferrer"
       style="display:inline-block;padding:0.5rem 1.2rem;background:#4a9eff;color:#fff;border-radius:4px;text-decoration:none;font-weight:600;">
      Open →
    </a>
  </div>
</div>"""
    except Exception as e:
        return (
            f'<div class="preview-error">Preview error: {html_lib.escape(str(e))}</div>'
        )
